"""Offline tests for the read-only stage snapshot workbook."""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from udise_vps import snapshot  # noqa: E402


class FakeSession:
    school_id = "2497128"
    students = [
        {"studentId": "11", "studentCodeNat": "PEN11", "studentName": "Asha", "classId": 9},
        {"studentId": "12", "studentCodeNat": "PEN12", "studentName": "Ravi", "classId": 10},
    ]

    def student_detail(self, sid):
        if sid == "12":
            raise RuntimeError("GP unavailable")
        return {
            "studentId": sid, "studentName": "Asha", "classDesc": "IX",
            "uuid": "123412341234", "formStatus": 3, "tokenValue": "must-not-leak",
        }

    def enrolment_detail(self, sid):
        return {"studentId": sid, "schoolId": self.school_id, "admnNumber": "7"}

    def facility_detail(self, sid):
        if sid == "12":
            raise RuntimeError("Facility unavailable")
        return {"studentId": sid, "schoolId": self.school_id, "heightInCm": 152}


def test_collect_and_write_snapshot():
    data = snapshot.collect_snapshot(FakeSession())
    assert len(data["Students"]) == 2
    assert len(data["GP"]) == 1
    assert len(data["EP"]) == 2
    assert len(data["Facility"]) == 1
    assert len(data["Completion"]) == 2
    assert len(data["Issues"]) == 2
    assert data["GP"][0]["uuid"] == "XXXX-XXXX-1234"
    assert data["GP"][0]["tokenValue"] == "[redacted]"

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "snapshot.xlsx"
        snapshot.write_snapshot_workbook(data, str(path))
        wb = load_workbook(path, read_only=True)
        assert wb.sheetnames == ["Students", "GP", "EP", "Facility", "Completion", "Issues"]
        gp_rows = list(wb["GP"].iter_rows(values_only=True))
        flat = [value for row in gp_rows for value in row]
        assert "123412341234" not in flat
        assert "must-not-leak" not in flat
        assert "XXXX-XXXX-1234" in flat
        wb.close()


def test_gp_pending_uses_field_eligibility_not_form_status():
    class Live:
        students = [{"studentId": "11", "studentCodeNat": "PEN11", "studentName": "Asha", "classId": 9}]

        def student_detail(self, sid):
            from udise_vps.constants import AUTO_GP_DEFAULTS
            # Overall formStatus=0 must not make GP pending when every AUTO
            # field is already populated.
            return {"studentId": sid, "studentCodeNat": "PEN11", "studentName": "Asha",
                    "classId": 9, "formStatus": 0, "aayBplYN": 1, **{k: (2 if k == "cwsnYN" else 1) for k in AUTO_GP_DEFAULTS}}

        def enrolment_detail(self, sid): return {}
        def facility_detail(self, sid): return {}

    data = snapshot.collect_snapshot(Live(), "IX")
    assert data["Completion"][0]["formStatus"] == 0
    status, changes = snapshot.build_auto_gp_changes(data["GP"][0], "PEN11")
    assert status == "NO_CHANGE"
    assert changes == {}


def test_empty_roster_is_rejected():
    class Empty:
        students = []

    try:
        snapshot.collect_snapshot(Empty())
        raise AssertionError("expected RuntimeError")
    except RuntimeError as exc:
        assert "Roster" in str(exc)


if __name__ == "__main__":
    tests = [test_collect_and_write_snapshot, test_empty_roster_is_rejected]
    for test in tests:
        test()
        print(f"PASS {test.__name__}")
    print(f"{len(tests)}/{len(tests)} passed")
