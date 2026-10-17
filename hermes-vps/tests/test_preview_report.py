"""Offline tests for write-stage preview workbooks."""
import tempfile
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from udise_vps import preview_report  # noqa: E402


def test_preview_workbook_contains_proposed_values():
    results = [SimpleNamespace(
        pen="PEN1", student_id="11", name="Asha", class_label="IX",
        status="PREVIEW", detail="2 blank fields would be filled",
        changes={"motherTongue": 42, "bloodGroup": "9"},
    )]
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "gp-preview.xlsx"
        preview_report.write_preview_workbook("gp", results, str(path))
        wb = load_workbook(path, read_only=True)
        assert wb.sheetnames == ["Summary", "Proposed Changes"]
        rows = list(wb["Proposed Changes"].iter_rows(values_only=True))
        flat = [value for row in rows for value in row]
        assert "motherTongue" in flat
        assert 42 in flat
        assert "bloodGroup" in flat
        wb.close()


def test_write_result_workbook_has_result_mode_and_filename():
    results = [SimpleNamespace(
        pen="PEN2", student_id="22", name="Ravi", class_label="XI",
        status="SUCCESS_CONFIRMED", detail="Saved and verified by fresh read-back",
        changes={"motherTongue": 42},
    )]
    with tempfile.TemporaryDirectory() as directory:
        filename = preview_report.default_result_filename("gp", "SCHOOL1", "XI")
        assert filename.startswith("UDISE_GP_Save_Result_SCHOOL1_XI_")
        path = Path(directory) / filename
        preview_report.write_preview_workbook(
            "gp", results, str(path),
            report_mode="WRITE RESULTS - portal outcomes and verification status",
            details_sheet="Student Results",
        )
        wb = load_workbook(path, read_only=True)
        assert wb.sheetnames == ["Summary", "Student Results"]
        summary = list(wb["Summary"].iter_rows(values_only=True))
        assert ("Mode", "WRITE RESULTS - portal outcomes and verification status") in summary
        rows = list(wb["Student Results"].iter_rows(values_only=True))
        flat = [value for row in rows for value in row]
        assert "SUCCESS_CONFIRMED" in flat
        assert "Saved and verified by fresh read-back" in flat
        wb.close()


def test_eshiksha_sheet_masks_sensitive_values():
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "ep-preview.xlsx"
        preview_report.write_preview_workbook(
            "ep", [], str(path),
            eshiksha_rows=[{
                "Student": "Asha", "Aadhaar": "123412341234",
                "password": "must-not-leak",
            }],
        )
        wb = load_workbook(path, read_only=True)
        rows = list(wb["eShikshaKosh Source"].iter_rows(values_only=True))
        flat = [value for row in rows for value in row]
        assert "123412341234" not in flat
        assert "must-not-leak" not in flat
        assert "XXXX-XXXX-1234" in flat
        assert "[redacted]" in flat
        wb.close()


if __name__ == "__main__":
    tests = [
        test_preview_workbook_contains_proposed_values,
        test_write_result_workbook_has_result_mode_and_filename,
        test_eshiksha_sheet_masks_sensitive_values,
    ]
    for test in tests:
        test()
        print(f"PASS {test.__name__}")
    print(f"{len(tests)}/{len(tests)} passed")
