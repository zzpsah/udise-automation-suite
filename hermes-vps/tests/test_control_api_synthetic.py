"""Synthetic preview/approval control-plane test with no portal access."""

from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook


def test_preview_lifecycle_and_write_lock() -> None:
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
        root = Path(directory)
        token_file = root / "api-token"
        token_file.write_text("synthetic-token\n", encoding="utf-8")
        os.environ["UDISE_CONTROL_STATE"] = str(root / "state")
        os.environ["UDISE_CONTROL_RUNTIME"] = str(root / "runtime")
        os.environ["UDISE_CONTROL_TOKEN_FILE"] = str(token_file)
        os.environ["UDISE_VPS_RUNNER"] = str(root / "synthetic-runner")

        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
        from control_api import app as api

        session_id = "synthetic_session_identifier_12345"
        api._json_write(
            api.SESSIONS / f"{session_id}.json",
            {
                "cookie": "JSESSIONID=synthetic; XSRF-TOKEN=synthetic",
                "created_at": int(time.time()),
                "expires_at": int(time.time()) + 300,
            },
        )

        observed: dict[str, object] = {"commands": []}

        class SyntheticProcess:
            def __init__(self, cmd, cwd, env, **_kwargs):
                observed["cmd"] = list(cmd)
                observed["commands"].append(list(cmd))
                observed["cookie"] = env.get("UDISE_COOKIE_HEADER")
                out_dir = Path(cmd[cmd.index("--out") + 1])
                result = out_dir / "synthetic-gp-preview.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Proposed Changes"
                sheet.append(["PEN", "Field", "Current", "Proposed"])
                sheet.append(["SYNTHETIC-1", "motherTongue", "", "42"])
                workbook.save(result)
                if "--plan-out" in cmd:
                    plan_path = Path(cmd[cmd.index("--plan-out") + 1])
                    plan_path.parent.mkdir(parents=True, exist_ok=True)
                    plan_path.write_text(
                        '{"SYNTHETIC-1":{"pen":"SYNTHETIC-1","student_id":"synthetic-1","changes":{"motherTongue":42}}}\n',
                        encoding="utf-8",
                    )
                self.stdout = iter(
                    [
                        "Authenticated synthetic session\n",
                        "Loaded 2 students\n",
                        f"REPORT_READY={result}\n",
                    ]
                )

            @staticmethod
            def wait() -> int:
                return 0

        original_popen = api.subprocess.Popen
        original_check = api._check_portal_session
        api.subprocess.Popen = SyntheticProcess
        api._check_portal_session = lambda sid, extend=False: api._load_session(sid)
        try:
            authorization = "Bearer synthetic-token"
            capabilities = api.capabilities()
            gp = next(stage for stage in capabilities["stages"] if stage["id"] == "gp")
            assert gp["preview_enabled"] is True

            try:
                api.create_job(
                    api.JobIn(
                        session_id=session_id,
                        school="2497128",
                        stage="gp",
                        class_name="IX",
                        preview=False,
                    ),
                    authorization=authorization,
                )
            except HTTPException as exc:
                assert exc.status_code == 409
            else:
                raise AssertionError("A non-preview write-stage job was accepted")

            auto_queued = api.create_job(
                api.JobIn(
                    session_id=session_id,
                    school="2497128",
                    stage="gp",
                    class_name="IX",
                    preview=True,
                    auto_save=True,
                        max_submissions=7,
                ),
                authorization=authorization,
            )
            deadline = time.time() + 5
            auto_state = None
            while time.time() < deadline:
                auto_state = api.get_job(auto_queued["job_id"], authorization=authorization)
                if auto_state["job"].get("auto_write_job_id"):
                    break
                time.sleep(0.02)
            assert auto_state is not None
            assert auto_state["job"]["status"] == "completed", auto_state
            auto_write_id = auto_state["job"]["auto_write_job_id"]
            assert auto_write_id
            deadline = time.time() + 5
            auto_write_state = None
            while time.time() < deadline:
                auto_write_state = api.get_job(auto_write_id, authorization=authorization)
                if auto_write_state["job"]["status"] in {"completed", "failed"}:
                    break
                time.sleep(0.02)
            assert auto_write_state is not None
            assert auto_write_state["job"]["status"] == "completed", auto_write_state
            assert auto_write_state["job"]["preview"] == 0
            assert auto_write_state["job"]["approved_from"] == auto_queued["job_id"]
            auto_write_cmd = observed["commands"][-1]
            assert "--submit" in auto_write_cmd
            assert auto_write_cmd[auto_write_cmd.index("--max") + 1] == "7"

            queued = api.create_job(
                api.JobIn(
                    session_id=session_id,
                    school="2497128",
                    stage="gp",
                    class_name="IX",
                    preview=True,
                ),
                authorization=authorization,
            )
            deadline = time.time() + 5
            state = None
            while time.time() < deadline:
                state = api.get_job(queued["job_id"], authorization=authorization)
                if state["job"]["status"] in {"completed", "failed"}:
                    break
                time.sleep(0.02)

            assert state is not None
            assert state["job"]["status"] == "completed", state
            assert state["job"]["has_result"] is True
            assert state["job"]["progress_current"] == 2
            assert state["job"]["progress_total"] == 2
            assert "--submit" not in observed["cmd"]
            assert observed["cookie"] == "JSESSIONID=synthetic; XSRF-TOKEN=synthetic"

            with api._db() as connection:
                row = connection.execute(
                    "SELECT result_path FROM jobs WHERE id=?", (queued["job_id"],)
                ).fetchone()
            workbook = load_workbook(row["result_path"], read_only=True)
            assert workbook.sheetnames == ["Proposed Changes"]
            assert workbook["Proposed Changes"].max_row == 2
            workbook.close()

            try:
                api.approve_job(
                    queued["job_id"],
                    api.ApprovalIn(
                        confirmation="SAVE GP IX",
                        acknowledge_readback=False,
                        max_submissions=1,
                    ),
                    authorization=authorization,
                )
            except HTTPException as exc:
                assert exc.status_code == 400
            else:
                raise AssertionError("Approval without read-back acknowledgement was accepted")

            approved = api.approve_job(
                queued["job_id"],
                api.ApprovalIn(
                    confirmation="SAVE GP IX",
                    acknowledge_readback=True,
                    max_submissions=1,
                ),
                authorization=authorization,
            )
            deadline = time.time() + 5
            write_state = None
            while time.time() < deadline:
                write_state = api.get_job(approved["job_id"], authorization=authorization)
                if write_state["job"]["status"] in {"completed", "failed"}:
                    break
                time.sleep(0.02)
            assert write_state is not None
            assert write_state["job"]["status"] == "completed", write_state
            assert write_state["job"]["preview"] == 0
            assert write_state["job"]["approved_from"] == queued["job_id"]
            write_cmd = observed["commands"][-1]
            assert "--submit" in write_cmd
            assert write_cmd[write_cmd.index("--max") + 1] == "1"

            try:
                api.approve_job(
                    queued["job_id"],
                    api.ApprovalIn(
                        confirmation="SAVE GP IX",
                        acknowledge_readback=True,
                        max_submissions=1,
                    ),
                    authorization=authorization,
                )
            except HTTPException as exc:
                assert exc.status_code == 409
            else:
                raise AssertionError("The same preview was approved twice")

            ep_preview_id = "a" * 32
            ep_dir = api.JOBS / ep_preview_id
            ep_dir.mkdir(parents=True)
            ep_result = ep_dir / "UDISE_EP_Preview_2497128_IX.xlsx"
            ep_source = ep_dir / "eShikshaKosh_OTR_2497128_2026-27.xlsx"
            (ep_dir / "approved-plan.json").write_text(
                '{"SYNTHETIC-1":{"pen":"SYNTHETIC-1","student_id":"synthetic-1","changes":{"admnNumber":"1/2026"}}}\n',
                encoding="utf-8",
            )
            for path in (ep_result, ep_source):
                book = Workbook()
                book.save(path)
                book.close()
            now = int(time.time())
            with api._db() as connection:
                connection.execute(
                    """INSERT INTO jobs(
                           id,created_at,updated_at,status,stage,class_name,school,
                           session_id,preview,message,result_path
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        ep_preview_id, now, now, "completed", "ep", "IX", "2497128",
                        session_id, 1, "Completed successfully", str(ep_result),
                    ),
                )
            ep_approved = api.approve_job(
                ep_preview_id,
                api.ApprovalIn(
                    confirmation="SAVE EP IX",
                    acknowledge_readback=True,
                    max_submissions=5,
                ),
                authorization=authorization,
            )
            deadline = time.time() + 5
            while time.time() < deadline:
                ep_state = api.get_job(ep_approved["job_id"], authorization=authorization)
                if ep_state["job"]["status"] in {"completed", "failed"}:
                    break
                time.sleep(0.02)
            assert ep_state["job"]["status"] == "completed", ep_state
            ep_cmd = observed["commands"][-1]
            assert "--report" in ep_cmd and str(ep_source) in ep_cmd
            assert "--fetch-report" not in ep_cmd
            assert ep_cmd[ep_cmd.index("--max") + 1] == "5"

            for requested_class in ("IX", "X", "XI", "XII"):
                snapshot_job = api.create_job(
                    api.JobIn(
                        session_id=session_id,
                        school="2497128",
                        stage="snapshot",
                        class_name=requested_class,
                        preview=True,
                    ),
                    authorization=authorization,
                )
                deadline = time.time() + 5
                snapshot_state = None
                while time.time() < deadline:
                    snapshot_state = api.get_job(snapshot_job["job_id"], authorization=authorization)
                    if snapshot_state["job"]["status"] in {"completed", "failed"}:
                        break
                    time.sleep(0.02)
                assert snapshot_state is not None
                assert snapshot_state["job"]["status"] == "completed", snapshot_state
                snapshot_cmd = observed["commands"][-1]
                assert snapshot_cmd[snapshot_cmd.index("--class") + 1] == requested_class
                assert "--submit" not in snapshot_cmd
        finally:
            api.subprocess.Popen = original_popen
            api._check_portal_session = original_check


if __name__ == "__main__":
    test_preview_lifecycle_and_write_lock()
    print("PASS synthetic preview, approval, write cap, and duplicate lock")
