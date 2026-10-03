"""Class Completion Overview — read-only status scan.

Ported from the notebook's Completion Overview cell. Groups students by the
project-observed formStatus progression 0 -> 1 -> 2 -> 3 -> 6 and exposes only
status 3 to Finalize.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from .constants import CLASS_LABEL, STATUS_ORDER, STATUS_STAGE, class_scope

STATUS_GUIDANCE = {
    0: "🔴 GP + EP + FP pending",
    1: "🟡 GP completed — complete EP + FP",
    2: "🟠 GP + EP completed — complete FP",
    3: "🟢 GP + EP + FP completed — READY TO COMPLETE DATA",
    6: "✅ Complete Data completed",
}


@dataclass
class CompletionRow:
    pen: str
    student_id: str
    name: str
    class_label: str
    form_status: Any
    stage: str
    guidance: str
    profile_status: Any = None
    status_desc: Any = None
    last_modified: Any = None


@dataclass
class CompletionReport:
    rows: list = field(default_factory=list)
    failed: list = field(default_factory=list)
    ready_pens: list = field(default_factory=list)
    completed_pens: list = field(default_factory=list)
    all_pending_pens: list = field(default_factory=list)
    gp_only_pens: list = field(default_factory=list)
    gp_ep_pens: list = field(default_factory=list)
    unknown_pens: list = field(default_factory=list)
    scope: str = ""


def scan_completion(
    session,
    *,
    class_scope_name: str = "IX",
    get_attempts: int = 2,
) -> CompletionReport:
    """Read every selected student's current status. Performs no writes."""
    classes = class_scope(class_scope_name)

    selected = [
        s for s in session.students
        if int(s.get("classId", -1)) in classes
    ]
    if not selected:
        raise ValueError(
            f"No students found in scope {class_scope_name}. "
            "Re-run login if the roster is stale."
        )

    report = CompletionReport(scope=class_scope_name)
    total = len(selected)
    print(f"[COMPLETION] Checking {total} students in {class_scope_name}.", flush=True)

    for position, student in enumerate(selected, 1):
        sid = str(student.get("studentId") or student.get("id") or "").strip()
        pen = str(student.get("studentCodeNat") or "").strip()
        class_label = CLASS_LABEL.get(int(student.get("classId", 0)), "")

        try:
            data = session.student_detail(sid)
            raw = data.get("formStatus")
            try:
                form_status = int(raw)
            except (TypeError, ValueError):
                form_status = None

            row = CompletionRow(
                pen=pen,
                student_id=sid,
                name=data.get("studentName") or student.get("studentName", ""),
                class_label=class_label,
                form_status=form_status,
                stage=STATUS_STAGE.get(form_status, "Unknown"),
                guidance=STATUS_GUIDANCE.get(form_status, "⚠️ Unknown status — manual review"),
                profile_status=data.get("profileStatus"),
                status_desc=data.get("statusDesc"),
                last_modified=data.get("lastModifiedOn"),
            )
            report.rows.append(row)
            print(
                f"[COMPLETION] {position}/{total} {pen}: OK "
                f"(formStatus={form_status}, {row.stage})",
                flush=True,
            )

        except Exception as exc:
            report.failed.append(pen or sid)
            report.rows.append(CompletionRow(
                pen=pen, student_id=sid,
                name=student.get("studentName", ""),
                class_label=class_label,
                form_status="ERROR", stage="Error",
                guidance="⚠️ Could not read status",
                status_desc=str(exc),
            ))
            print(
                f"[COMPLETION] {position}/{total} {pen}: ⚠️ "
                f"{type(exc).__name__}: {exc}",
                flush=True,
            )

    def pens_for(status: int) -> list:
        return [r.pen for r in report.rows if r.form_status == status and r.pen]

    report.all_pending_pens = pens_for(0)
    report.gp_only_pens = pens_for(1)
    report.gp_ep_pens = pens_for(2)
    report.ready_pens = pens_for(3)
    report.completed_pens = pens_for(6)
    report.unknown_pens = [
        r.pen for r in report.rows
        if r.form_status not in {0, 1, 2, 3, 6, "ERROR"} and r.pen
    ]

    print("\n━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    print(f"✅ Completed             : {len(report.completed_pens)}")
    print(f"🟢 Ready to Complete     : {len(report.ready_pens)}")
    print(f"🟠 Need FP               : {len(report.gp_ep_pens)}")
    print(f"🟡 Need EP + FP          : {len(report.gp_only_pens)}")
    print(f"🔴 Need GP + EP + FP     : {len(report.all_pending_pens)}")
    if report.unknown_pens:
        print(f"⚠️ Unknown status        : {len(report.unknown_pens)}")
    if report.failed:
        print(f"❌ Read failures         : {len(report.failed)}")
    print("━━━━━━━━━━━━━━━━━━━━━━━━━━━━")

    if report.ready_pens:
        print(
            f"\n🟢 {len(report.ready_pens)} student(s) are ready for "
            "Complete Data / Finalize."
        )
        print("AUTO Finalize will use only these (formStatus=3).")
    else:
        print("\nℹ️ No formStatus=3 students are currently ready.")

    return report


def report_to_rows(report: CompletionReport) -> list[dict]:
    """Flatten the report for Excel/CSV export."""
    rows = [
        {
            "PEN": r.pen,
            "Student ID (system)": r.student_id,
            "Name": r.name,
            "Class": r.class_label,
            "formStatus": r.form_status,
            "Stage": r.stage,
            "Guidance": r.guidance,
            "profileStatus": r.profile_status,
            "statusDesc": r.status_desc,
            "lastModifiedOn": r.last_modified,
        }
        for r in report.rows
    ]
    rows.sort(key=lambda r: (
        STATUS_ORDER.get(r["formStatus"], 99) if isinstance(r["formStatus"], int) else 99,
        str(r["Name"]),
    ))
    return rows


def write_completion_workbook(report: CompletionReport, path: str) -> str:
    """Write the multi-sheet overview workbook, one sheet per status group."""
    from openpyxl import Workbook

    rows = report_to_rows(report)
    wb = Workbook()

    groups = [
        ("Overview", rows),
        ("Need GP EP FP", [r for r in rows if r["formStatus"] == 0]),
        ("Need EP FP", [r for r in rows if r["formStatus"] == 1]),
        ("Need FP", [r for r in rows if r["formStatus"] == 2]),
        ("Ready to Complete", [r for r in rows if r["formStatus"] == 3]),
        ("Completed", [r for r in rows if r["formStatus"] == 6]),
        ("Errors Unknown", [r for r in rows if r["formStatus"] not in {0, 1, 2, 3, 6}]),
    ]

    first = True
    for title, group in groups:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        first = False
        if group:
            headers = list(group[0].keys())
            ws.append(headers)
            for row in group:
                ws.append([row.get(h) for h in headers])
            for idx, _ in enumerate(headers, 1):
                ws.column_dimensions[ws.cell(1, idx).column_letter].width = 20

    wb.save(path)
    return path


def default_filename(school_id: str, scope: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"UDISE_Completion_Overview_{school_id}_{scope.replace(' ', '')}_{stamp}.xlsx"
