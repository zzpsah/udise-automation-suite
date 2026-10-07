"""Excel reports for read-only previews of write stages."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from .snapshot import _safe_value


def _result_rows(results: list[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        common = {
            "PEN": getattr(result, "pen", ""),
            "Student ID (system)": getattr(result, "student_id", ""),
            "Name": getattr(result, "name", ""),
            "Class": getattr(result, "class_label", ""),
            "Status": getattr(result, "status", ""),
            "Reason": getattr(result, "detail", "") or getattr(result, "reason", ""),
        }
        changes = getattr(result, "changes", {}) or {}
        proposed = getattr(result, "proposed", {}) or {}
        if isinstance(changes, list):
            changes = {field: proposed.get(field, "") for field in changes}
        if isinstance(changes, dict) and changes:
            for field, value in changes.items():
                rows.append({**common, "Field": field, "Proposed Value": _safe_value(field, value)})
        elif getattr(result, "status", "") == "PREVIEW" and result.__class__.__name__ == "FinalizeResult":
            rows.append({**common, "Field": "formStatus", "Proposed Value": 6})
        else:
            rows.append({**common, "Field": "", "Proposed Value": ""})
    return rows


def _write_rows(ws, rows: list[dict[str, Any]], headers: list[str] | None = None) -> None:
    if headers is None:
        headers = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
    headers = headers or ["Status"]
    fill = PatternFill("solid", fgColor="1F4E79")
    for column, header in enumerate(headers, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            ws.cell(row_index, column, _safe_value(header, row.get(header, "")))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, header in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, column).column_letter].width = min(42, max(12, len(header) + 2))


def write_preview_workbook(
    stage: str,
    results: list[Any],
    path: str,
    *,
    eshiksha_rows: list[dict] | None = None,
) -> str:
    """Write a no-save preview workbook, optionally with masked eShiksha rows."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    counts = Counter(str(getattr(result, "status", "UNKNOWN")) for result in results)
    summary_rows = [
        {"Item": "Stage", "Value": stage.upper()},
        {"Item": "Mode", "Value": "PREVIEW ONLY - NO PORTAL WRITE"},
        {"Item": "Students checked", "Value": len(results)},
        {"Item": "Generated", "Value": datetime.now().isoformat(timespec="seconds")},
    ] + [{"Item": f"Status: {status}", "Value": count} for status, count in sorted(counts.items())]
    _write_rows(summary, summary_rows)

    proposed = wb.create_sheet("Proposed Changes")
    _write_rows(proposed, _result_rows(results), [
        "PEN", "Student ID (system)", "Name", "Class", "Status",
        "Field", "Proposed Value", "Reason",
    ])

    if stage.lower() == "ep":
        review = wb.create_sheet("Enrollment Profile Review")
        ep_fields = [
            "admnNumber", "admnStartDate", "rollNumber", "moiId", "academicStream",
            "enrStatusPY", "classPY", "rteQuestion", "rteAmount", "examResultPy",
            "examMarksPy", "attendancePy",
        ] + [f"subject{i}" for i in range(1, 9)]
        review_headers = ["PEN", "Name", "Class", "Status"]
        for field in ep_fields:
            review_headers += [f"{field} Current", f"{field} Proposed", f"{field} Effective"]
        review_rows = []
        for result in results:
            current = getattr(result, "current", {}) or {}
            changes = getattr(result, "changes", {}) or {}
            row = {
                "PEN": getattr(result, "pen", ""),
                "Name": getattr(result, "name", ""),
                "Class": getattr(result, "class_label", ""),
                "Status": getattr(result, "status", ""),
            }
            for field in ep_fields:
                cur = current.get(field, "")
                proposed_value = changes.get(field, "")
                effective = proposed_value if field in changes else cur
                row[f"{field} Current"] = cur
                row[f"{field} Proposed"] = proposed_value
                row[f"{field} Effective"] = effective
            review_rows.append(row)
        _write_rows(review, review_rows, review_headers)

        lists = wb.create_sheet("Subject Lists")
        fallback = {
            1: ["HINDI", "URDU"],
            2: ["SANSKRIT", "HIN (NLH)"],
            3: ["MATHEMATICS"],
            4: ["SCIENCE"],
            5: ["SOCIAL SCIENCE"],
            6: ["ENGLISH"],
        }
        for slot in range(1, 7):
            source_col = slot
            lists.cell(1, source_col, f"Subject {slot}")
            for ri, label in enumerate(fallback[slot], 2):
                lists.cell(ri, source_col, label)
            target_header = f"subject{slot} Effective"
            target_col = review_headers.index(target_header) + 1
            target_letter = review.cell(1, target_col).column_letter
            source_letter = lists.cell(1, source_col).column_letter
            dv = DataValidation(
                type="list",
                formula1=f"'Subject Lists'!${source_letter}$2:${source_letter}${1 + len(fallback[slot])}",
                allow_blank=True,
            )
            review.add_data_validation(dv)
            if review_rows:
                dv.add(f"{target_letter}2:{target_letter}{len(review_rows) + 1}")
        lists.sheet_state = "hidden"

    if eshiksha_rows is not None:
        source = wb.create_sheet("eShikshaKosh Source")
        safe_rows = [
            {str(key): _safe_value(str(key), value) for key, value in row.items()}
            for row in eshiksha_rows
        ]
        _write_rows(source, safe_rows, ["No source rows"] if not safe_rows else None)

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return str(target)


def default_filename(stage: str, school_id: str, scope: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    clean_scope = scope.replace(" ", "")
    return f"UDISE_{stage.upper()}_Preview_{school_id}_{clean_scope}_{stamp}.xlsx"
