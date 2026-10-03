"""Read-only, stage-wise UDISE workbook export.

The snapshot performs only the same authenticated GETs already used by the
runner. It writes one workbook with Students, GP, EP, Facility, Completion and
Issues sheets. Aadhaar-like values are masked and secret-like fields are never
written to the workbook.
"""

from __future__ import annotations

import json
from collections.abc import Mapping
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from .constants import CLASS_LABEL, STATUS_STAGE
from .students import get_apaar_id, mask_aadhaar

_SECRET_PARTS = ("password", "token", "cookie", "secret", "authorization", "jsession", "xsrf")
_AADHAAR_PARTS = ("aadhaar", "aadhar", "uuid")


def _excel_text(value: str) -> str:
    """Remove illegal controls and prevent external text becoming a formula."""
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    return "'" + cleaned if cleaned.startswith(("=", "+", "-", "@")) else cleaned


def _safe_value(key: str, value: Any) -> Any:
    """Return an Excel-safe value without exposing secrets or full Aadhaar."""
    lowered = key.lower()
    if any(part in lowered for part in _SECRET_PARTS):
        return "[redacted]"
    if any(part in lowered for part in _AADHAAR_PARTS):
        return mask_aadhaar(value)
    if value is None:
        return ""
    if isinstance(value, str):
        return _excel_text(value)
    if isinstance(value, (int, float, bool)):
        return value
    return _excel_text(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str))


def _flatten(data: Mapping[str, Any], prefix: str = "") -> dict[str, Any]:
    row: dict[str, Any] = {}
    for key, value in data.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, Mapping):
            row.update(_flatten(value, name))
        else:
            row[name] = _safe_value(name, value)
    return row


def _identity(student: Mapping[str, Any], position: int) -> dict[str, Any]:
    class_id = student.get("classId")
    try:
        class_label = CLASS_LABEL.get(int(class_id), str(class_id or ""))
    except (TypeError, ValueError):
        class_label = str(class_id or "")
    return {
        "S. No.": position,
        "PEN": str(student.get("studentCodeNat") or ""),
        "Student ID (system)": str(student.get("studentId") or student.get("id") or ""),
        "Name": student.get("studentName") or "",
        "Class": class_label,
    }


def _student_summary(identity: dict[str, Any], roster: Mapping[str, Any], gp: Mapping[str, Any] | None) -> dict[str, Any]:
    source = gp or roster
    return {
        **identity,
        "Section": source.get("sectionDesc") or roster.get("sectionDesc") or "",
        "Mother Name": source.get("motherName") or "",
        "Father Name": source.get("fatherName") or "",
        "Date of Birth": source.get("dob") or "",
        "Mobile No.": source.get("primaryMobile") or "",
        "Masked Aadhaar": mask_aadhaar(source.get("uuid")),
        "APAAR ID": get_apaar_id(dict(source)),
    }


def collect_snapshot(session) -> dict[str, list[dict[str, Any]]]:
    """Read all stage records for the loaded roster. Performs no writes."""
    if not session.students:
        raise RuntimeError("Roster is not loaded. Run login first.")

    result: dict[str, list[dict[str, Any]]] = {
        "Students": [], "GP": [], "EP": [], "Facility": [],
        "Completion": [], "Issues": [],
    }
    total = len(session.students)
    print(f"[SNAPSHOT] Reading GP, EP and Facility for {total} students.", flush=True)

    for position, student in enumerate(session.students, 1):
        identity = _identity(student, position)
        sid = str(identity["Student ID (system)"])
        gp: dict[str, Any] | None = None

        try:
            gp = session.student_detail(sid)
            result["GP"].append({**identity, **_flatten(gp)})
        except Exception as exc:
            result["Issues"].append({**identity, "Stage": "GP", "Error": f"{type(exc).__name__}: {exc}"})

        try:
            ep = session.enrolment_detail(sid)
            result["EP"].append({**identity, **_flatten(ep)})
        except Exception as exc:
            result["Issues"].append({**identity, "Stage": "EP", "Error": f"{type(exc).__name__}: {exc}"})

        try:
            facility = session.facility_detail(sid)
            result["Facility"].append({**identity, **_flatten(facility)})
        except Exception as exc:
            result["Issues"].append({**identity, "Stage": "Facility", "Error": f"{type(exc).__name__}: {exc}"})

        result["Students"].append(_student_summary(identity, student, gp))
        raw_status = gp.get("formStatus") if gp else None
        try:
            form_status: Any = int(raw_status)
        except (TypeError, ValueError):
            form_status = "ERROR" if gp is None else raw_status
        result["Completion"].append({
            **identity,
            "formStatus": form_status,
            "Stage": STATUS_STAGE.get(form_status, "Unknown" if gp else "GP read failed"),
            "profileStatus": gp.get("profileStatus", "") if gp else "",
            "statusDesc": gp.get("statusDesc", "") if gp else "",
            "lastModifiedOn": gp.get("lastModifiedOn", "") if gp else "",
        })
        print(f"[SNAPSHOT] {position}/{total} students checked", flush=True)

    return result


def _write_sheet(ws, rows: list[dict[str, Any]], fallback_headers: list[str]) -> None:
    headers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                headers.append(key)
    if not headers:
        headers = fallback_headers

    for column, header in enumerate(headers, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            ws.cell(row_index, column, row.get(header, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, header in enumerate(headers, 1):
        sample = [len(str(row.get(header, ""))) for row in rows[:100]]
        ws.column_dimensions[ws.cell(1, column).column_letter].width = min(42, max(12, len(header) + 2, *(sample or [0])))


def write_snapshot_workbook(snapshot: dict[str, list[dict[str, Any]]], path: str) -> str:
    """Write the stage-wise snapshot workbook."""
    wb = Workbook()
    fallback = {
        "Students": ["S. No.", "PEN", "Student ID (system)", "Name", "Class"],
        "GP": ["S. No.", "PEN", "Student ID (system)", "Name", "Class"],
        "EP": ["S. No.", "PEN", "Student ID (system)", "Name", "Class"],
        "Facility": ["S. No.", "PEN", "Student ID (system)", "Name", "Class"],
        "Completion": ["S. No.", "PEN", "Student ID (system)", "Name", "Class", "formStatus", "Stage"],
        "Issues": ["S. No.", "PEN", "Student ID (system)", "Name", "Class", "Stage", "Error"],
    }
    for index, title in enumerate(("Students", "GP", "EP", "Facility", "Completion", "Issues")):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        _write_sheet(ws, snapshot.get(title, []), fallback[title])
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return str(target)


def default_filename(school_id: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"UDISE_Full_Read_Snapshot_{school_id}_{stamp}.xlsx"
