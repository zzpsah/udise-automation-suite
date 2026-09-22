# %%
#@title 🏫 Facility Profile — Choose Class { display-mode: "form" }
import io
import time
import threading
from datetime import datetime
from decimal import Decimal, InvalidOperation
import pandas as pd
import requests
from google.colab import files

FACILITY_CLASS = "IX" #@param ["IX", "X", "IX and X"]
FACILITY_BUILD = "v1.2.8 (2026-09-23)"
FACILITY_CLASSES = {"IX": {9}, "X": {10}, "IX and X": {9, 10}}[FACILITY_CLASS]
globals().pop('facility_reviewed', None)
FACILITY_BENEFITS = dict(enumerate(['Free Text Book', 'Free Uniforms', 'Free Transport facility', 'Free Bi-Cycle', 'Free hostel', 'Free Escort', 'Free Mobile/Tablet/Computer', 'Other'], 1))
FACILITY_CWSN = dict(enumerate(['Braille Book', 'Braille Kit', 'Braces', 'Tri-cycle', 'Stipend', 'Crutches', 'Caliper', 'Low Vision Kit', 'Hearing Aid', 'Wheel Chair', 'Escort', 'Other'], 1))
FACILITY_DISTANCE = {1: 'Less than 1 km', 2: 'Between 1-3 Kms', 3: 'Between 3-5 Kms', 4: 'More than 5 Kms'}
FACILITY_EDUCATION = {1: 'Primary', 2: 'Upper Primary', 3: 'Secondary or Equivalent', 4: 'Higher Secondary or Equivalent', 5: 'More than Higher Secondary', 6: 'No Schooling Experience'}
FACILITY_YN = {'Facilities Provided': 'facilityYn', 'CWSN Facilities Provided': 'facProvidedCwsnYn', 'Competitions/Olympiads': 'olympdsNlc', 'NCC': 'nccYn', 'NSS': 'nssYn', 'Scouts and Guides': 'scoutsYn'}

def facility_text(value):
    return '' if value is None or str(value).strip().lower() in {'nan', 'none', '<na>'} else str(value).strip()

def facility_int(value, label, low, high):
    try:
        num = Decimal(facility_text(value))
        if not num.is_finite() or num != num.to_integral_value() or not low <= num <= high:
            raise ValueError()
        return int(num)
    except (InvalidOperation, ValueError):
        raise ValueError(f'{label}: enter a whole number from {low} to {high}')

def facility_yes_no(value, label):
    text = facility_text(value).lower()
    if text in {'yes', '1'}:
        return 1
    if text in {'no', '2'}:
        return 2
    raise ValueError(f'{label}: choose Yes or No; blank/9 is unanswered')

def facility_label(value, mapping):
    if value is None or str(value) in {'0', '9', ''}:
        return ''
    code = facility_int(value, 'Code', 1, 99)
    return f'{code} - {mapping[code]}' if code in mapping else f'UNKNOWN CODE {code}'

def facility_choice(value, mapping, label):
    text = facility_text(value)
    choices = {f'{code} - {name}': code for code, name in mapping.items()}
    if text in choices:
        return choices[text]
    if text.isdigit() and int(text) in mapping:
        return int(text)
    raise ValueError(f'{label}: select a value from the dropdown')

def facility_request(method, route, **kwargs):
    if 'session' not in globals() or 'HEADERS' not in globals():
        raise RuntimeError('Run common Authentication first')
    done = threading.Event()
    start = time.monotonic()
    def heartbeat():
        while not done.wait(10):
            print(f'[FACILITY] WAIT {method}: {time.monotonic()-start:.0f}s', flush=True)
    worker = threading.Thread(target=heartbeat, daemon=True)
    worker.start()
    try:
        response = session.request(method, BASE_URL + route, headers=HEADERS, timeout=(15, 300 if method == 'POST' else 60), allow_redirects=False, **kwargs)
        print(f'[FACILITY] {method} HTTP {response.status_code} after {time.monotonic()-start:.1f}s', flush=True)
        try:
            body = response.json()
        except ValueError:
            body = {}
        return response.status_code, body if isinstance(body, dict) else {}
    finally:
        done.set()
        worker.join(timeout=1)

def facility_get(sid):
    status, body = facility_request('GET', f'/p0/api/v2/students/facility/{sid}')
    if status != 200 or body.get('status') is not True or not isinstance(body.get('data'), dict):
        raise RuntimeError(f'Facility record unavailable: HTTP {status}')
    data = body['data']
    if str(data.get('schoolId')) != str(SCHOOL_ID) or str(data.get('studentId')) != str(sid):
        raise ValueError('Facility record school/student identity mismatch')
    return data

def facility_roster():
    if 'students' not in globals() or 'SCHOOL_ID' not in globals():
        raise RuntimeError('Run Detect School and Fetch Current Students first')
    return {str(s['studentId']): s for s in students if int(s.get('classId') or -1) in FACILITY_CLASSES}

def facility_payload(row, cwsn):
    payload = {'schoolId': int(SCHOOL_ID)}
    for label, field in FACILITY_YN.items():
        if field == 'facProvidedCwsnYn' and cwsn is None:
            continue  # Applicability is checked after the live General Profile read.
        if field == 'facProvidedCwsnYn' and not cwsn:
            if facility_text(row.get(label)).lower() not in {'', 'na', 'not applicable'}:
                raise ValueError('CWSN Facilities Provided must be blank for a non-CWSN student')
            payload[field] = 9
        else:
            payload[field] = facility_yes_no(row.get(label), label)
    for prefix, mapping, flag, field in [('Benefit: ', FACILITY_BENEFITS, 'facilityYn', 'facProvided'), ('CWSN: ', FACILITY_CWSN, 'facProvidedCwsnYn', 'facProvidedCwsn')]:
        if flag == 'facProvidedCwsnYn' and cwsn is None:
            continue
        selected = []
        for code, label in mapping.items():
            raw = facility_text(row.get(prefix + label))
            if raw and facility_yes_no(raw, prefix + label) == 1:
                selected.append(code)
        if payload[flag] == 1 and not selected:
            raise ValueError(f'{prefix}choose at least one facility when Yes')
        if payload[flag] != 1 and selected:
            raise ValueError(f'{prefix}items cannot be Yes when the parent field is No/not applicable')
        payload[field] = selected if payload[flag] == 1 else None
    payload['heightInCm'] = str(facility_int(row.get('Height (cm)'), 'Height (cm)', 60, 256))
    payload['weightInKg'] = str(facility_int(row.get('Weight (kg)'), 'Weight (kg)', 10, 150))
    payload['distanceFrmSchool'] = str(facility_choice(row.get('Distance to School'), FACILITY_DISTANCE, 'Distance to School'))
    payload['parentEducation'] = str(facility_choice(row.get('Parent/Guardian Education'), FACILITY_EDUCATION, 'Parent/Guardian Education'))
    return payload

def facility_mismatches(saved, payload):
    def normal(key, value):
        if key in {'facProvided', 'facProvidedCwsn'}:
            return sorted(int(v) for v in (value or []))
        return str(value) if value is not None else ''
    return [key for key, value in payload.items() if normal(key, saved.get(key)) != normal(key, value)]

print(f'✅ Class {FACILITY_CLASS} selected. Facility choices are ready.')
# %%
#@title 📥 Facility Profile — Download Excel { display-mode: "form" }
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from tqdm.auto import tqdm

FACILITY_FETCH_LIMIT = 0 #@param {type:"integer"}
roster = facility_roster()
selected = list(roster.items())[:FACILITY_FETCH_LIMIT or None]
if not selected:
    raise ValueError('No students found in the selected class')
rows, failed = [], []
for sid, student in tqdm(selected, desc='Facility profiles'):
    try:
        data = facility_get(sid)
        status, general = facility_request('GET', f'/p0/api/cy/students/{sid}')
        cwsn_flag = (general.get('data') or {}).get('cwsnYN')
        if status != 200 or general.get('status') is not True or cwsn_flag not in (1, 2):
            raise ValueError('Cannot establish current CWSN status; profile not exported')
        row = {'PEN': str(student.get('studentCodeNat', '')), 'Student Name': student.get('studentName', ''), 'Class': {9: 'IX', 10: 'X'}[int(student['classId'])], 'Student ID (system)': sid}
        row['CWSN Student (reference)'] = 'Yes' if cwsn_flag == 1 else 'No'
        for label, field in FACILITY_YN.items():
            # Keep a saved Yes, but make unanswered choices ready as No.
            value = facility_text(data.get(field))
            if value not in {'1', '2', '9', '0', ''}:
                raise ValueError(f'Unknown {field} value; profile not exported')
            row[label] = 'Yes' if value == '1' else 'No'
        if cwsn_flag == 2:
            row['CWSN Facilities Provided'] = ''
        for prefix, mapping, field in [('Benefit: ', FACILITY_BENEFITS, 'facProvided'), ('CWSN: ', FACILITY_CWSN, 'facProvidedCwsn')]:
            codes = {int(v) for v in (data.get(field) or [])}
            if codes - set(mapping):
                raise ValueError('Unknown facility option codes: refresh portal reference data')
            row.update({prefix + label: 'Yes' if code in codes else 'No' for code, label in mapping.items()})
        if cwsn_flag == 2:
            for label in FACILITY_CWSN.values():
                row['CWSN: '+label] = 'No'
        row.update({'Height (cm)': data.get('heightInCm') or '', 'Weight (kg)': data.get('weightInKg') or '', 'Distance to School': facility_label(data.get('distanceFrmSchool'), FACILITY_DISTANCE) or facility_label(2, FACILITY_DISTANCE), 'Parent/Guardian Education': facility_label(data.get('parentEducation'), FACILITY_EDUCATION) or facility_label(3, FACILITY_EDUCATION)})
        rows.append(row)
    except Exception as exc:
        failed.append({'Student ID (system)': sid, 'Error': str(exc)})
        print(f'[FACILITY] EXPORT ERROR: {type(exc).__name__}', flush=True)
if not rows:
    raise RuntimeError('No profiles exported; check authentication and response logs')
book = Workbook()
sheet = book.active
sheet.title = 'Facility Update'
columns = list(rows[0])
sheet.append(columns)
for row in rows:
    sheet.append([row[c] for c in columns])
sheet.freeze_panes = 'F2'
sheet.auto_filter.ref = sheet.dimensions
for c in sheet[1]:
    c.fill = PatternFill('solid', fgColor='1F4E78')
    c.font = Font(color='FFFFFF', bold=True)
for i, col in enumerate(columns, 1):
    sheet.column_dimensions[get_column_letter(i)].width = min(38, max(18, len(col)+2))
    for row_number in range(2, sheet.max_row+1):
        cell = sheet.cell(row_number, i)
        non_cwsn = rows[row_number-2]['CWSN Student (reference)'] == 'No'
        locked = i <= 5 or (non_cwsn and (col == 'CWSN Facilities Provided' or col.startswith('CWSN: ')))
        cell.protection = Protection(locked=locked)
        if locked:
            cell.fill = PatternFill('solid', fgColor='E7E6E6')
sheet.protection.sheet = True
sheet.protection.autoFilter = False
for label, low, high in [('Height (cm)', 60, 256), ('Weight (kg)', 10, 150)]:
    column = get_column_letter(columns.index(label)+1)
    rule = DataValidation(type='whole', operator='between', formula1=str(low), formula2=str(high), allow_blank=True)
    rule.showErrorMessage = True
    rule.errorStyle = 'stop'
    rule.error = f'Enter the actual measured value, whole numbers {low}–{high}.'
    rule.showInputMessage = True
    rule.prompt = f'Actual measurement required ({low}–{high}); do not estimate by gender.'
    sheet.add_data_validation(rule)
    rule.add(f'{column}2:{column}{sheet.max_row}')
lists = book.create_sheet('Dropdown Lists')
lists.sheet_state = 'hidden'
options = {label: ['Yes', 'No'] for label in columns[5:] if label not in {'Height (cm)', 'Weight (kg)', 'Distance to School', 'Parent/Guardian Education'}}
options['Distance to School'] = [f'{k} - {v}' for k, v in FACILITY_DISTANCE.items()]
options['Parent/Guardian Education'] = [f'{k} - {v}' for k, v in FACILITY_EDUCATION.items()]
off_col = get_column_letter(len(options)+1)
lists.cell(1, len(options)+1, 'No')
book.defined_names.add(DefinedName('facility_disabled', attr_text=f"'Dropdown Lists'!${off_col}$1"))
for i, (label, choices) in enumerate(options.items(), 1):
    letter = get_column_letter(i)
    for j, value in enumerate(choices, 1):
        lists.cell(j, i, value)
    name = f'facility_options_{i}'
    book.defined_names.add(DefinedName(name, attr_text=f"'Dropdown Lists'!${letter}$1:${letter}${len(choices)}"))
    column = get_column_letter(columns.index(label)+1)
    formula = '='+name
    parent = 'Facilities Provided' if label.startswith('Benefit: ') else 'CWSN Facilities Provided' if label.startswith('CWSN: ') else None
    if parent:
        parent_col = get_column_letter(columns.index(parent)+1)
        formula = f'IF(${parent_col}2="Yes",{name},facility_disabled)'
        sheet.conditional_formatting.add(f'{column}2:{column}{sheet.max_row}', FormulaRule(formula=[f'${parent_col}2<>"Yes"'], fill=PatternFill('solid', fgColor='E7E6E6'), font=Font(color='808080')))
        sheet.conditional_formatting.add(f'{column}2:{column}{sheet.max_row}', FormulaRule(formula=[f'AND(${parent_col}2<>"Yes",{column}2="Yes")'], fill=PatternFill('solid', fgColor='FFC7CE'), stopIfTrue=True))
    validation = DataValidation(type='list', formula1=formula, allow_blank=True)
    validation.errorStyle = 'stop'
    validation.showErrorMessage = True
    validation.error = 'Choose a listed value. Benefits require the parent field to be Yes.'
    sheet.add_data_validation(validation)
    validation.add(f'{column}2:{column}{sheet.max_row}')
if failed:
    error_sheet = book.create_sheet('Export Errors')
    error_sheet.append(['Student ID (system)', 'Error'])
    for error in failed:
        error_sheet.append(list(error.values()))
facility_export_file = f'UDISE_Facility_{FACILITY_CLASS.replace(" ", "_")}_{SCHOOL_ID}_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
book.save(facility_export_file)
print(f'📥 Workbook ready: {len(rows)} students. {len(failed)} could not be included. Saved measurements are shown; enter actual height or weight only where blank.')
files.download(facility_export_file)
# %%
#@title ✅ Facility Profile — Check Uploaded Excel { display-mode: "form" }
globals().pop('facility_reviewed', None)
upload = files.upload()
if len(upload) != 1:
    raise ValueError('Upload exactly one Facility Update workbook')
facility_frame = pd.read_excel(io.BytesIO(next(iter(upload.values()))), sheet_name='Facility Update', dtype=object)
facility_frame = facility_frame.where(pd.notna(facility_frame), '')
required = {'PEN', 'Class', 'Student ID (system)', 'CWSN Student (reference)', 'Height (cm)', 'Weight (kg)', 'Distance to School', 'Parent/Guardian Education'} | set(FACILITY_YN) | {'Benefit: '+v for v in FACILITY_BENEFITS.values()} | {'CWSN: '+v for v in FACILITY_CWSN.values()}
if required - set(facility_frame):
    raise ValueError('Missing columns: ' + ', '.join(sorted(required-set(facility_frame))))
roster = facility_roster()
reviewed, issues, seen = [], [], set()
print('🔎 Checking the Excel file. No student record will be changed.', flush=True)
for index, row in facility_frame.iterrows():
    sid = facility_text(row['Student ID (system)'])
    try:
        if sid in seen or sid not in roster:
            raise ValueError('Duplicate student or student outside selected roster')
        seen.add(sid)
        student = roster[sid]
        if facility_text(row['PEN']) != str(student.get('studentCodeNat', '')) or row['Class'] != {9: 'IX', 10: 'X'}[int(student['classId'])]:
            raise ValueError('PEN/class does not match current roster')
        reference = facility_text(row['CWSN Student (reference)']).lower()
        if reference not in {'yes', 'no'}:
            raise ValueError('CWSN Student (reference) must be Yes or No. Export a fresh Facility workbook.')
        cwsn_flag = 1 if reference == 'yes' else 2
        payload = facility_payload(row, cwsn_flag == 1)
        reviewed.append({'sid': sid, 'pen': facility_text(row['PEN']), 'cwsn': cwsn_flag, 'payload': payload})
    except (ValueError, TypeError, KeyError) as exc:
        issues.append((index+2, str(exc)))
        print(f'✏️ Excel row {index+2} needs attention: {exc}', flush=True)
if issues or not reviewed:
    print(f'⚠️ Please fix {len(issues)} Excel issue(s) and check the file again. Nothing was submitted.')
else:
    facility_reviewed = {'school': str(SCHOOL_ID), 'class': FACILITY_CLASS, 'rows': reviewed}
    print(f'✅ Excel check passed for {len(reviewed)} student(s). Nothing was submitted. Review the file before choosing Submit.')
# %%
#@title 🚀 Facility Profile — Submit Reviewed Updates { display-mode: "form" }
ALLOW_FACILITY_UPDATE = False #@param {type:"boolean"}
FACILITY_MAX_SUBMISSIONS = 1 #@param {type:"integer"}
if not ALLOW_FACILITY_UPDATE:
    raise RuntimeError('Enable ALLOW_FACILITY_UPDATE after reviewing and validating the workbook')
if 'facility_reviewed' not in globals() or facility_reviewed['school'] != str(SCHOOL_ID) or facility_reviewed['class'] != FACILITY_CLASS:
    raise RuntimeError('Run Facility validation for the current school/class first')
if FACILITY_MAX_SUBMISSIONS < 1:
    raise ValueError('Submission limit must be at least one')
facility_results = []
result_file = f'UDISE_Facility_Result_{SCHOOL_ID}_v1.2.8_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
for position, item in enumerate(facility_reviewed['rows'][:FACILITY_MAX_SUBMISSIONS], 1):
    sid, payload = item['sid'], item['payload']
    result = {'PEN': item['pen'], 'Student ID (system)': sid, 'Status': 'FAILED', 'Detail': '', 'Build': FACILITY_BUILD}
    try:
        print(f'🔎 Student {position}: checking the current portal record', flush=True)
        current = facility_get(sid)
        changes = facility_mismatches(current, payload)
        result['Changed Fields'] = ', '.join(changes)
        if not changes:
            result.update(Status='SKIPPED_ALREADY_UP_TO_DATE', Detail='Fresh record already matches')
        else:
            status, general = facility_request('GET', f'/p0/api/cy/students/{sid}')
            if status != 200 or general.get('status') is not True or (general.get('data') or {}).get('cwsnYN') != item['cwsn']:
                raise ValueError('CWSN status changed or unavailable; validate again')
            result.update(Status='SUBMISSION_PENDING', Detail='Check fresh portal state before any retry')
            pd.DataFrame(facility_results+[result]).to_excel(result_file, index=False)
            try:
                status, body = facility_request('POST', f'/p0/api/v2/AY/students/facility/{sid}', json=payload)
                error = body.get('error') or {}
                detail = error.get('message') if isinstance(error, dict) else str(error)
                result['Detail'] = f'HTTP {status}; {detail or body.get("message") or "No response message"}'
                if isinstance(error, dict):
                    fields = (error.get('data') or {}).get('errorFields')
                    if fields:
                        result['Detail'] += '; ' + str(fields)
            except requests.RequestException as exc:
                result['Detail'] = f'POST {type(exc).__name__}; outcome unknown'
            print('📨 Portal response received; checking whether the changes were saved.', flush=True)
            result['Status'] = 'UNCONFIRMED'
            for attempt, delay in enumerate((2, 5, 10), 1):
                print(f'🔄 Confirming saved details ({attempt}/3)', flush=True)
                time.sleep(delay)
                try:
                    changes = facility_mismatches(facility_get(sid), payload)
                    if not changes:
                        result.update(Status='SUCCESS_CONFIRMED_BY_READBACK', Detail=result['Detail']+'; saved fields confirmed')
                        break
                    result.update(Status='FAILED', Detail=result['Detail']+'; mismatched: '+', '.join(changes))
                except Exception as exc:
                    print('⚠️ Could not confirm the saved details yet:', type(exc).__name__, flush=True)
    except Exception as exc:
        result['Detail'] += '; ' + str(exc)
    friendly_status = {'SKIPPED_ALREADY_UP_TO_DATE': '✅ Already up to date — no change sent', 'SUCCESS_CONFIRMED_BY_READBACK': '✅ Saved and confirmed', 'FAILED': '❌ Not saved — check the result file', 'UNCONFIRMED': '⚠️ Save not confirmed — check the portal before retrying'}.get(result['Status'], '⚠️ Check the result file')
    result['Result for user'] = friendly_status
    facility_results.append(result)
    pd.DataFrame(facility_results).to_excel(result_file, index=False)
    print(f'Student {position}: {friendly_status}', flush=True)
    if result['Status'] not in {'SKIPPED_ALREADY_UP_TO_DATE', 'SUCCESS_CONFIRMED_BY_READBACK'}:
        break
files.download(result_file)
