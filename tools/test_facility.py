import contextlib
import io
import sys
import types
import unittest
import tempfile
import os
from pathlib import Path
from unittest.mock import patch
import pandas as pd
import requests
import openpyxl

class FacilityTests(unittest.TestCase):
    def setUp(self):
        source = Path(__file__).with_name('facility_cells.py').read_text(encoding='utf-8').split('# %%\n')[1]
        colab = types.ModuleType('google.colab')
        colab.files = types.SimpleNamespace()
        self.env = {'SCHOOL_ID': '123'}
        with patch.dict(sys.modules, {'google.colab': colab}), contextlib.redirect_stdout(io.StringIO()):
            exec(source, self.env)
        self.row = {label: 'No' for label in self.env['FACILITY_YN']}
        self.row.update({'CWSN Facilities Provided': '', 'Height (cm)': 150, 'Weight (kg)': 40, 'Distance to School': '1 - Less than 1 km', 'Parent/Guardian Education': '4 - Higher Secondary or Equivalent'})

    def test_validation_local_first_and_error_classes(self):
        source = Path(__file__).with_name('facility_cells.py').read_text(encoding='utf-8').split('# %%\n')[3]
        for case in ('input', 'success', 'missing_measurement', 'missing_reference'):
            with self.subTest(case=case):
                row = dict(self.row, PEN='p1', Class='IX', **{'Student ID (system)': 's1', 'CWSN Student (reference)': 'No'})
                for values, prefix in ((self.env['FACILITY_BENEFITS'], 'Benefit: '), (self.env['FACILITY_CWSN'], 'CWSN: ')):
                    row.update({prefix+v: 'No' for v in values.values()})
                if case == 'input':
                    row['NCC'] = 9
                if case == 'missing_measurement':
                    row['Height (cm)'] = ''
                if case == 'missing_reference':
                    row.pop('CWSN Student (reference)')
                calls = []
                def request(*args):
                    calls.append(args)
                    raise AssertionError('Sheet validation must not call UDISE')
                self.env.update(facility_reviewed='stale', files=types.SimpleNamespace(upload=lambda: {'test.xlsx': b'x'}),
                                facility_roster=lambda: {'s1': {'studentCodeNat': 'p1', 'classId': 9}}, facility_request=request)
                output = io.StringIO()
                with patch.object(pd, 'read_excel', return_value=pd.DataFrame([row])), contextlib.redirect_stdout(output):
                    if case == 'missing_reference':
                        with self.assertRaisesRegex(ValueError, 'CWSN Student'):
                            exec(source, self.env)
                    else:
                        exec(source, self.env)
                self.assertEqual(len(calls), 0)
                self.assertEqual('facility_reviewed' in self.env, case == 'success')
                if case == 'input':
                    self.assertIn('Excel row 2 needs attention', output.getvalue())
                if case == 'missing_measurement':
                    self.assertIn('Height (cm)', output.getvalue())

    def test_payload_fields_and_non_cwsn(self):
        p = self.env['facility_payload'](self.row, False)
        self.assertEqual(set(p), {'schoolId','facilityYn','facProvided','facProvidedCwsnYn','facProvidedCwsn','olympdsNlc','nccYn','nssYn','scoutsYn','heightInCm','weightInKg','distanceFrmSchool','parentEducation'})
        self.assertEqual(p['facProvidedCwsnYn'], 9)
        self.assertIsNone(p['facProvided'])
        self.assertEqual(p['nccYn'], 2)

    def test_unanswered_is_not_no(self):
        self.row['NCC'] = '9'
        with self.assertRaisesRegex(ValueError, 'unanswered'):
            self.env['facility_payload'](self.row, False)

    def test_yes_requires_detail(self):
        self.row['Facilities Provided'] = 'Yes'
        with self.assertRaisesRegex(ValueError, 'at least one'):
            self.env['facility_payload'](self.row, False)
        self.row['Benefit: Free Text Book'] = 'Yes'
        self.assertEqual(self.env['facility_payload'](self.row, False)['facProvided'], [1])

    def test_non_cwsn_rejects_support(self):
        self.row['CWSN: Braille Book'] = 'Yes'
        with self.assertRaises(ValueError):
            self.env['facility_payload'](self.row, False)

    def test_ranges_and_decimal(self):
        for v in [0, 59, 257, '150.5']:
            self.row['Height (cm)'] = v
            with self.assertRaises(ValueError):
                self.env['facility_payload'](self.row, False)

    def test_readback_type_normalization(self):
        p = self.env['facility_payload'](self.row, False)
        saved = {**p, 'heightInCm': 150, 'weightInKg': 40, 'distanceFrmSchool': 1, 'parentEducation': 4}
        self.assertEqual(self.env['facility_mismatches'](saved, p), [])
        saved['nccYn'] = 1
        self.assertEqual(self.env['facility_mismatches'](saved, p), ['nccYn'])

    def test_submission_scenarios(self):
        source = Path(__file__).with_name('facility_cells.py').read_text(encoding='utf-8').split('# %%\n')[-1].replace('ALLOW_FACILITY_UPDATE = False', 'ALLOW_FACILITY_UPDATE = True')
        for case in ('success', 'timeout_saved', 'rejected'):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as folder:
                payload = self.env['facility_payload'](self.row, False)
                before = {**payload, 'nccYn': 9}
                reads = [before, payload] if case != 'rejected' else [before]*4
                calls = []
                def request(method, route, **kwargs):
                    calls.append((method, route))
                    if method == 'GET':
                        return 200, {'status': True, 'data': {'cwsnYN': 2}}
                    self.assertEqual(route, '/p0/api/v2/AY/students/facility/test')
                    if case == 'timeout_saved':
                        raise requests.Timeout()
                    return 200, {'status': case == 'success', 'message': 'Successful' if case == 'success' else 'Rejected'}
                self.env.update(facility_reviewed={'school': '123', 'class': 'IX', 'rows': [{'sid': 'test', 'pen': 'test', 'payload': payload, 'cwsn': 2}]}, facility_get=lambda sid: reads.pop(0), facility_request=request, files=types.SimpleNamespace(download=lambda p: None))
                cwd = os.getcwd()
                try:
                    os.chdir(folder)
                    with patch('time.sleep'), contextlib.redirect_stdout(io.StringIO()):
                        exec(source, self.env)
                    expected = 'FAILED' if case == 'rejected' else 'SUCCESS_CONFIRMED_BY_READBACK'
                    self.assertEqual(self.env['facility_results'][0]['Status'], expected)
                    self.assertEqual(sum(method == 'POST' for method, _ in calls), 1)
                    self.assertTrue(list(Path(folder).glob('*.xlsx')))
                finally:
                    os.chdir(cwd)

    def test_export_conditional_controls(self):
        source = Path(__file__).with_name('facility_cells.py').read_text(encoding='utf-8').split('# %%\n')[2]
        colab = types.ModuleType('google.colab')
        colab.files = types.SimpleNamespace(download=lambda p: None)
        auto = types.ModuleType('tqdm.auto')
        auto.tqdm = lambda sequence, **kwargs: sequence
        data = self.env['facility_payload'](self.row, False)
        data.update(nccYn=9, nssYn=9, olympdsNlc=1, distanceFrmSchool=9)
        missing_education = dict(data, parentEducation=9)
        self.env.update(students=[{'studentId': 's1', 'studentCodeNat': 'p1', 'classId': 9}, {'studentId': 's2', 'studentCodeNat': 'p2', 'classId': 10}], FACILITY_CLASSES={9, 10}, facility_get=lambda sid: missing_education if sid == 's1' else data, facility_request=lambda *args: (200, {'status': True, 'data': {'cwsnYN': 2}}), files=colab.files)
        with tempfile.TemporaryDirectory() as folder:
            cwd = os.getcwd()
            try:
                os.chdir(folder)
                output = io.StringIO()
                with patch.dict(sys.modules, {'google.colab': colab, 'tqdm.auto': auto}), contextlib.redirect_stdout(output):
                    exec(source, self.env)
                self.assertIn('50% (1/2)', output.getvalue())
                self.assertIn('100% (2/2)', output.getvalue())
                book = openpyxl.load_workbook(self.env['facility_export_file'])
                sheet = book['Facility Update']
                headers = {c.value: c.column for c in sheet[1]}
                self.assertTrue(sheet.protection.sheet)
                self.assertEqual(sheet.cell(2, headers['CWSN Student (reference)']).value, 'No')
                self.assertEqual(sheet.cell(2, headers['Facilities Provided']).value, 'No')
                self.assertEqual(sheet.cell(2, headers['NCC']).value, 'No')
                self.assertEqual(sheet.cell(2, headers['NSS']).value, 'No')
                self.assertEqual(sheet.cell(2, headers['Competitions/Olympiads']).value, 'Yes')
                self.assertEqual(sheet.cell(2, headers['Distance to School']).value, '2 - Between 1-3 Kms')
                self.assertEqual(sheet.cell(2, headers['Parent/Guardian Education']).value, '3 - Secondary or Equivalent')
                self.assertEqual(sheet.cell(3, headers['Parent/Guardian Education']).value, '4 - Higher Secondary or Equivalent')
                self.assertEqual(sheet.cell(2, headers['Height (cm)']).value, '150')
                self.assertEqual(sheet.cell(2, headers['Weight (kg)']).value, '40')
                for field in ('CWSN Facilities Provided', 'CWSN: Braille Book'):
                    self.assertTrue(sheet.cell(2, headers[field]).protection.locked)
                self.assertFalse(sheet.cell(2, headers['Facilities Provided']).protection.locked)
                rules = list(sheet.data_validations.dataValidation)
                self.assertTrue(any(str(r.formula1).startswith('IF(') and 'facility_disabled' in r.formula1 for r in rules))
                self.assertTrue(any(r.type == 'whole' and r.formula1 == '60' and r.formula2 == '256' for r in rules))
                self.assertTrue(any(r.type == 'whole' and r.formula1 == '10' and r.formula2 == '150' for r in rules))
                book.close()
            finally:
                os.chdir(cwd)

if __name__ == '__main__':
    unittest.main()
